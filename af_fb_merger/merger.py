"""
AppsFlyer x Facebook data merger.

Pure data logic — no UI dependencies. Callable from gui.py or a script.

Key behaviors:
  - AF exports are often at ad level (Campaign × Ad). We auto-detect this by
    the presence of an "ad" column and aggregate up to campaign.
  - FB exports typically include a total/summary row with an empty campaign
    name — we drop it.
  - Metric ownership follows user rules: spend/impressions/clicks/registrations
    come from FB; installs/logins/paying_users/revenue come from AF.
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from column_aliases import (
    AF_ALIASES,
    CAMPAIGN_OUTPUT_COLUMNS,
    COUNTRY_CAMPAIGN_OUTPUT_COLUMNS,
    DISPLAY_HEADERS,
    FB_ALIASES,
    REGION_OUTPUT_COLUMNS,
)

UNMATCHED_FILL = PatternFill(start_color="FFD7D7", end_color="FFD7D7", fill_type="solid")
HEADER_FILL = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
TOTALS_FILL = PatternFill(start_color="FFF4CC", end_color="FFF4CC", fill_type="solid")
SUBTOTAL_FILL = PatternFill(start_color="E7F0F9", end_color="E7F0F9", fill_type="solid")
REGION_TOTAL_FILL = PatternFill(start_color="D5E3F0", end_color="D5E3F0", fill_type="solid")
HEADER_FONT = Font(bold=True)
TOTALS_FONT = Font(bold=True)

_ENCODING_CANDIDATES = ("utf-8-sig", "utf-8", "gbk", "gb18030", "cp936")

# Numeric canonicals that should be summed when aggregating AF ad→campaign.
_AF_SUM_COLS = ("installs", "logins", "paying_users", "purchase_count", "revenue")
_AF_RATE_COLS = (
    "login_rate", "pay_rate",
    "arppu", "revenue_per_purchase",
    "retention_d1", "retention_d2", "retention_d6", "retention_d7",
    "retention_d14", "retention_d21", "retention_d30",
)
# Numeric canonicals in FB that need to be coerced from string to float.
_FB_NUMERIC = ("spend", "impressions", "clicks", "registrations")


class ColumnNotFoundError(Exception):
    """Raised when a required canonical column cannot be mapped from the raw file."""


def _normalize(s: str) -> str:
    return "".join(str(s).split()).lower()


def _resolve_columns(
    raw_columns: List[str],
    aliases: Dict[str, List[str]],
    required: List[str],
) -> Dict[str, str]:
    """Return {canonical_name: raw_column_name} for each canonical key that matches.

    Raises ColumnNotFoundError if any key in `required` is missing.
    """
    normalized = {_normalize(c): c for c in raw_columns}
    resolved: Dict[str, str] = {}
    for canonical, candidates in aliases.items():
        for cand in candidates:
            key = _normalize(cand)
            if key in normalized:
                resolved[canonical] = normalized[key]
                break
    missing = [k for k in required if k not in resolved]
    if missing:
        raise ColumnNotFoundError(
            f"Missing required columns: {missing}.\n"
            f"Found raw columns: {raw_columns}.\n"
            f"Please add the actual column name to column_aliases.py."
        )
    return resolved


def _read_tabular(path: str | Path) -> pd.DataFrame:
    """Read CSV (auto-encoding) or XLSX into a DataFrame."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xls"):
        return pd.read_excel(path)
    last_err: Optional[Exception] = None
    for enc in _ENCODING_CANDIDATES:
        try:
            return pd.read_csv(path, encoding=enc)
        except UnicodeDecodeError as e:
            last_err = e
            continue
    raise UnicodeDecodeError(
        "utf-8", b"", 0, 1,
        f"Could not decode {path} with any of {_ENCODING_CANDIDATES}: {last_err}"
    )


def _coerce_numeric(df: pd.DataFrame, cols) -> None:
    for c in cols:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce")


def load_af(path: str | Path, *, level: str = "campaign") -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Load AppsFlyer export; aggregate to campaign level if ad-level data detected.

    Returns (canonical_df, raw_df). `level` is 'campaign', 'region', or
    'country_campaign' (requires both country and campaign columns).
    """
    raw = _read_tabular(path)
    if level == "country_campaign":
        required_keys = ["campaign", "country"]
    elif level == "region":
        required_keys = ["country"]
    else:
        required_keys = ["campaign"]
    resolved = _resolve_columns(list(raw.columns), AF_ALIASES, required=required_keys)

    df = raw.rename(columns={v: k for k, v in resolved.items()})
    df = df[[c for c in resolved.keys() if c in df.columns]].copy()

    for key in required_keys:
        if key in df.columns:
            df[key] = df[key].astype(str).str.strip()
            df = df[df[key].str.len() > 0]
            df = df[df[key].str.lower() != "nan"]

    _coerce_numeric(df, _AF_SUM_COLS)
    _coerce_numeric(df, _AF_RATE_COLS)

    # Ad-level detection: "ad" column present with ≥1 non-empty value → aggregate.
    if level in ("campaign", "country_campaign") and "ad" in df.columns and df["ad"].notna().any():
        group_cols = ["campaign"]
        if level == "country_campaign" and "country" in df.columns:
            group_cols = ["country", "campaign"]
        sum_cols = [c for c in _AF_SUM_COLS if c in df.columns]
        agg_map = {c: "sum" for c in sum_cols}
        for c in _AF_RATE_COLS:
            if c in df.columns:
                agg_map[c] = "mean"
        df = df.groupby(group_cols, as_index=False).agg(agg_map)
        if {"logins", "installs"}.issubset(df.columns):
            df["login_rate"] = _safe_div(df["logins"], df["installs"])
        if {"paying_users", "installs"}.issubset(df.columns):
            df["pay_rate"] = _safe_div(df["paying_users"], df["installs"])
        if {"revenue", "paying_users"}.issubset(df.columns):
            df["arppu"] = _safe_div(df["revenue"], df["paying_users"])
        if {"revenue", "purchase_count"}.issubset(df.columns):
            df["revenue_per_purchase"] = _safe_div(df["revenue"], df["purchase_count"])

    return df, raw


def load_fb(path: str | Path, *, level: str = "campaign") -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Load Facebook Ads Manager export; drop the summary total row."""
    raw = _read_tabular(path)
    if level == "country_campaign":
        required_keys = ["campaign", "country"]
    elif level == "region":
        required_keys = ["country"]
    else:
        required_keys = ["campaign"]
    resolved = _resolve_columns(list(raw.columns), FB_ALIASES, required=required_keys + ["spend"])

    df = raw.rename(columns={v: k for k, v in resolved.items()})
    df = df[[c for c in resolved.keys() if c in df.columns]].copy()

    for key in required_keys:
        if key in df.columns:
            df[key] = df[key].astype(str).str.strip()
            df = df[df[key].str.len() > 0]
            df = df[df[key].str.lower() != "nan"]

    _coerce_numeric(df, _FB_NUMERIC)
    return df, raw


def _safe_div(num: pd.Series, den: pd.Series) -> pd.Series:
    return num.div(den.where(den != 0)).astype(float)


def merge_campaign(af: pd.DataFrame, fb: pd.DataFrame) -> pd.DataFrame:
    """Outer join AF and FB on campaign and compute derived metrics."""
    merged = af.merge(fb, on="campaign", how="outer", indicator=True)
    merged["_match_status"] = merged["_merge"].map({
        "both": "matched", "left_only": "af_only", "right_only": "fb_only"
    })
    merged = merged.drop(columns=["_merge"])

    _add_derived_metrics(merged)

    cols = [c for c in CAMPAIGN_OUTPUT_COLUMNS if c in merged.columns]
    out = merged[cols].copy()
    if "spend" in out.columns:
        out = out.sort_values("spend", ascending=False, na_position="last")
    out = out.reset_index(drop=True)
    totals = _build_totals_row(out, key_col="campaign")
    return pd.concat([out, pd.DataFrame([totals])], ignore_index=True)


def parse_campaign_dims(name: str) -> Tuple[str, str]:
    """Extract (OS, Goal) from a campaign name.

    OS detection looks for 'iOS' or 'Android' as a hyphen-separated token.
    Goal detection prefers primary objectives (MAI/AEO); if neither is present,
    falls back to variant labels (Purchase/Retargeting). Anything unrecognized
    maps to 'Other'.

    Example inputs (token split by '-'):
      ROG-01-FB-SEA-Android-PrimeCreative-260417-MAI-Retargeting → (Android, MAI)
      ROG-FB-SEA-Android-AIimage-260211-#zitou-Purchase          → (Android, Purchase)
    """
    if not isinstance(name, str) or not name:
        return ("Other", "Other")
    tokens = [t.upper() for t in name.split("-")]

    if "IOS" in tokens:
        os_ = "iOS"
    elif "ANDROID" in tokens:
        os_ = "Android"
    else:
        os_ = "Other"

    primary = {"MAI": "MAI", "AEO": "AEO"}
    for t in tokens:
        if t in primary:
            return (os_, primary[t])
    variant = {"PURCHASE": "Purchase", "RETARGETING": "Retargeting"}
    for t in tokens:
        if t in variant:
            return (os_, variant[t])
    return (os_, "Other")


def merge_country_campaign(af: pd.DataFrame, fb: pd.DataFrame) -> pd.DataFrame:
    """Outer join on (country, campaign) and annotate with OS/Goal."""
    merged = af.merge(fb, on=["country", "campaign"], how="outer", indicator=True)
    merged["_match_status"] = merged["_merge"].map({
        "both": "matched", "left_only": "af_only", "right_only": "fb_only"
    })
    merged = merged.drop(columns=["_merge"])
    _add_derived_metrics(merged)
    merged[["os", "goal"]] = merged["campaign"].apply(
        lambda c: pd.Series(parse_campaign_dims(c))
    )
    return merged


def build_country_pivot(merged: pd.DataFrame) -> pd.DataFrame:
    """Aggregate to OS × Country level. Structure:

      Android | MY | sum across all MY Android campaigns
      Android | PH | ...
      ...
      Android | — | Android 合计
      iOS | MY | ...
      ...
      iOS | — | iOS 合计
      总计

    Within each OS block, countries are sorted by spend desc. Order of OS
    blocks: Android first, then iOS, then anything else.
    """
    df = merged.copy()
    for c in ("country", "os"):
        if c in df.columns:
            df[c] = df[c].fillna("—")

    rows: List[Dict] = []

    def _row_from_group(group: pd.DataFrame, *, os_, country, match_status) -> Dict:
        row = _build_totals_row(group, key_col="campaign")
        row["os"] = os_
        row["country"] = country
        row["campaign_count"] = int(group["campaign"].nunique())
        row["_match_status"] = match_status
        row.pop("campaign", None)
        return row

    def _match_label(group: pd.DataFrame) -> str:
        statuses = set(group["_match_status"].dropna().unique())
        if statuses == {"matched"}:
            return "全匹配"
        if "af_only" in statuses and "fb_only" not in statuses:
            return "部分 af_only"
        if "fb_only" in statuses and "af_only" not in statuses:
            return "部分 fb_only"
        return "混合"

    os_order = ["Android", "iOS"]
    other_os = [o for o in df["os"].unique() if o not in os_order]
    ordered_os = os_order + sorted(other_os)

    for os_ in ordered_os:
        g_os = df[df["os"] == os_]
        if g_os.empty:
            continue
        # Per-country aggregation within this OS, sorted by spend desc.
        country_spend = g_os.groupby("country")["spend"].sum(min_count=1)
        country_order = country_spend.sort_values(ascending=False, na_position="last").index
        for country in country_order:
            g_c = g_os[g_os["country"] == country]
            rows.append(_row_from_group(
                g_c, os_=os_, country=country, match_status=_match_label(g_c),
            ))
        # OS subtotal
        rows.append(_row_from_group(
            g_os, os_=os_, country="— 合计", match_status="小计",
        ))

    # Grand total
    rows.append(_row_from_group(
        df, os_="—", country="总计", match_status="总计",
    ))

    out_cols = [c for c in COUNTRY_CAMPAIGN_OUTPUT_COLUMNS if c in rows[0]]
    return pd.DataFrame(rows)[out_cols]


def merge_region(af: pd.DataFrame, fb: pd.DataFrame) -> pd.DataFrame:
    """Outer join AF and FB on country."""
    merged = af.merge(fb, on="country", how="outer", indicator=True)
    merged["_match_status"] = merged["_merge"].map({
        "both": "matched", "left_only": "af_only", "right_only": "fb_only"
    })
    merged = merged.drop(columns=["_merge"])

    _add_derived_metrics(merged)

    cols = [c for c in REGION_OUTPUT_COLUMNS if c in merged.columns]
    out = merged[cols].copy()
    if "spend" in out.columns:
        out = out.sort_values("spend", ascending=False, na_position="last")
    out = out.reset_index(drop=True)
    totals = _build_totals_row(out, key_col="country")
    return pd.concat([out, pd.DataFrame([totals])], ignore_index=True)


def _add_derived_metrics(merged: pd.DataFrame) -> None:
    """Compute in-place: CTR, CPM, click→install CVR, CPI, CPA, ROAS."""
    if "clicks" in merged and "impressions" in merged:
        merged["ctr"] = _safe_div(merged["clicks"], merged["impressions"])
    if "spend" in merged and "impressions" in merged:
        merged["cpm"] = _safe_div(merged["spend"] * 1000, merged["impressions"])
    if "installs" in merged and "clicks" in merged:
        merged["click_to_install_cvr"] = _safe_div(merged["installs"], merged["clicks"])
    if "spend" in merged and "installs" in merged:
        merged["cpi"] = _safe_div(merged["spend"], merged["installs"])
    if "spend" in merged and "paying_users" in merged:
        merged["cpa"] = _safe_div(merged["spend"], merged["paying_users"])
    if "spend" in merged and "revenue" in merged:
        merged["roas"] = _safe_div(merged["revenue"], merged["spend"])


_SUMMABLE_COLS = (
    "spend", "impressions", "clicks",
    "installs", "logins", "paying_users", "purchase_count", "revenue",
    "registrations",
)
_RETENTION_COLS = (
    "retention_d1", "retention_d2", "retention_d6", "retention_d7",
    "retention_d14", "retention_d21", "retention_d30",
)


def _build_totals_row(df: pd.DataFrame, *, key_col: str) -> Dict:
    """Compute totals row: sums for raw metrics, recomputed rates from those sums,
    installs-weighted mean for retention rates. Non-derivable cells stay blank.
    """
    row: Dict = {c: None for c in df.columns}
    row[key_col] = "总计"
    row["_match_status"] = "总计"

    for c in _SUMMABLE_COLS:
        if c in df.columns:
            row[c] = df[c].sum(skipna=True)

    # installs-weighted mean for retention
    if "installs" in df.columns:
        weights = df["installs"].fillna(0)
        for c in _RETENTION_COLS:
            if c not in df.columns:
                continue
            mask = df[c].notna() & (weights > 0)
            if mask.any():
                w = weights[mask]
                v = df.loc[mask, c]
                total_w = w.sum()
                if total_w > 0:
                    row[c] = float((v * w).sum() / total_w)

    # Derived metrics recomputed from summed totals (not averaged per row).
    def _d(n, d):
        try:
            n = float(n) if n is not None else 0
            d = float(d) if d is not None else 0
        except (TypeError, ValueError):
            return None
        return n / d if d else None

    spend = row.get("spend")
    impressions = row.get("impressions")
    clicks = row.get("clicks")
    installs = row.get("installs")
    logins = row.get("logins")
    paying = row.get("paying_users")
    purchase_count = row.get("purchase_count")
    revenue = row.get("revenue")

    if "cpm" in df.columns:                   row["cpm"] = _d((spend or 0) * 1000, impressions)
    if "ctr" in df.columns:                   row["ctr"] = _d(clicks, impressions)
    if "click_to_install_cvr" in df.columns:  row["click_to_install_cvr"] = _d(installs, clicks)
    if "cpi" in df.columns:                   row["cpi"] = _d(spend, installs)
    if "login_rate" in df.columns:            row["login_rate"] = _d(logins, installs)
    if "cpa" in df.columns:                   row["cpa"] = _d(spend, paying)
    if "pay_rate" in df.columns:              row["pay_rate"] = _d(paying, installs)
    if "roas" in df.columns:                  row["roas"] = _d(revenue, spend)
    if "arppu" in df.columns:                 row["arppu"] = _d(revenue, paying)
    if "revenue_per_purchase" in df.columns:  row["revenue_per_purchase"] = _d(revenue, purchase_count)

    return row


_PERCENT_COLS = {
    "ctr", "click_to_install_cvr", "login_rate", "pay_rate",
    "retention_d1", "retention_d2", "retention_d6", "retention_d7",
    "retention_d14", "retention_d21", "retention_d30",
}
_MONEY_COLS = {"spend", "cpm", "cpi", "cpa", "revenue", "arppu", "revenue_per_purchase"}
_RATIO_COLS = {"roas"}
_INT_COLS = {
    "impressions", "clicks", "registrations",
    "installs", "logins", "paying_users", "purchase_count",
}


def _apply_number_formats(ws, columns: List[str]) -> None:
    """Apply XLSX number formats to columns based on canonical name."""
    for col_idx, col_name in enumerate(columns, start=1):
        fmt = None
        if col_name in _PERCENT_COLS:
            fmt = "0.00%"
        elif col_name in _MONEY_COLS:
            # USD with $ prefix, 2 decimals, thousand separator
            fmt = '"$"#,##0.00'
        elif col_name in _RATIO_COLS:
            fmt = "0.0000"
        elif col_name in _INT_COLS:
            fmt = "#,##0"
        if fmt is None:
            continue
        letter = get_column_letter(col_idx)
        for row in range(2, ws.max_row + 1):
            ws[f"{letter}{row}"].number_format = fmt


def _write_sheet(writer, sheet_name: str, df: pd.DataFrame, *, use_display_headers: bool) -> None:
    out = df.copy()
    canonical_cols = list(out.columns)
    if use_display_headers:
        out = out.rename(columns=DISPLAY_HEADERS)
    out.to_excel(writer, sheet_name=sheet_name, index=False)

    ws = writer.sheets[sheet_name]

    for col_idx in range(1, len(out.columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, col_name in enumerate(out.columns, start=1):
        max_len = max(
            [len(str(col_name))] + [len(str(v)) for v in out[col_name].astype(str).head(200)]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 32)

    if use_display_headers:
        _apply_number_formats(ws, canonical_cols)

    status_col_name = DISPLAY_HEADERS["_match_status"] if use_display_headers else "_match_status"
    if status_col_name in out.columns:
        status_col_idx = list(out.columns).index(status_col_name) + 1
        for row_idx in range(2, len(out) + 2):
            status = ws.cell(row=row_idx, column=status_col_idx).value
            fill = None
            bold = False
            if status in ("af_only", "fb_only"):
                fill = UNMATCHED_FILL
            elif status == "总计":
                fill = TOTALS_FILL
                bold = True
            elif status == "地区合计":
                fill = REGION_TOTAL_FILL
                bold = True
            elif status == "小计":
                fill = SUBTOTAL_FILL
                bold = True
            if fill is not None:
                for col_idx in range(1, len(out.columns) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.fill = fill
                    if bold:
                        cell.font = TOTALS_FONT

    ws.freeze_panes = "A2"


def export_xlsx(
    out_path: str | Path,
    *,
    campaign_merged: Optional[pd.DataFrame] = None,
    region_merged: Optional[pd.DataFrame] = None,
    country_campaign_pivot: Optional[pd.DataFrame] = None,
    raw_af_campaign: Optional[pd.DataFrame] = None,
    raw_fb_campaign: Optional[pd.DataFrame] = None,
    raw_af_region: Optional[pd.DataFrame] = None,
    raw_fb_region: Optional[pd.DataFrame] = None,
    raw_af_country_campaign: Optional[pd.DataFrame] = None,
    raw_fb_country_campaign: Optional[pd.DataFrame] = None,
) -> Path:
    """Write a single multi-sheet XLSX. Sheets are only added if data is provided."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        if campaign_merged is not None:
            _write_sheet(writer, "Campaign合并", campaign_merged, use_display_headers=True)
        if region_merged is not None:
            _write_sheet(writer, "地区合并", region_merged, use_display_headers=True)
        if country_campaign_pivot is not None:
            _write_sheet(writer, "国家×系列透视", country_campaign_pivot, use_display_headers=True)
        if raw_af_campaign is not None:
            _write_sheet(writer, "原始AF_Campaign", raw_af_campaign, use_display_headers=False)
        if raw_fb_campaign is not None:
            _write_sheet(writer, "原始FB_Campaign", raw_fb_campaign, use_display_headers=False)
        if raw_af_region is not None:
            _write_sheet(writer, "原始AF_地区", raw_af_region, use_display_headers=False)
        if raw_fb_region is not None:
            _write_sheet(writer, "原始FB_地区", raw_fb_region, use_display_headers=False)
        if raw_af_country_campaign is not None:
            _write_sheet(writer, "原始AF_国家系列", raw_af_country_campaign, use_display_headers=False)
        if raw_fb_country_campaign is not None:
            _write_sheet(writer, "原始FB_国家系列", raw_fb_country_campaign, use_display_headers=False)

    return out_path
